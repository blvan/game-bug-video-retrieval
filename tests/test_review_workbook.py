import csv
import json
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

try:
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font
    from scripts import review_workbook as review
except ImportError:
    review = None


@unittest.skipIf(review is None, 'Install requirements-review.txt to test the optional Excel workflow')
class ReviewLedgerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.run = self.root / 'pilot'
        self.run.mkdir()
        self.path = self.root / 'review.xlsx'
        self.a = {'repository': 'Game/Example', 'issue_number': '7',
                  'issue_title': '=untrusted text', 'issue_url': 'https://github.com/game/example/issues/7',
                  'video_url': 'https://example.org/a.mp4?key=1#part'}
        self.b = {**self.a, 'issue_number': '8', 'video_url': 'https://example.org/b.mp4'}
        self.write_run([self.a])

    def write_run(self, rows):
        with (self.run / 'video_issue_candidates.csv').open('w', newline='', encoding='utf-8') as stream:
            writer = csv.DictWriter(stream, fieldnames=list(self.a))
            writer.writeheader()
            writer.writerows(rows)

    def sync(self):
        return review.sync([self.run], self.path)

    def change(self, sheet, row, field, value):
        headers = [c.value for c in sheet[1]]
        sheet.cell(row, headers.index(field) + 1, value)

    def accept(self):
        book = load_workbook(self.path)
        for field in review.PAIR_CHECKS + ('manual_valid',):
            self.change(book['Pairs'], 2, field, 'yes')
        self.change(book['Pairs'], 2, 'bug_category', 'rendering')
        for field in review.REPO_CHECKS:
            self.change(book['Repositories'], 2, field, 'yes')
        book.save(self.path)
        book.close()

    def test_first_sync_is_unreviewed_and_text_is_literal(self):
        self.assertEqual(self.sync()['added_pairs'], 1)
        book = load_workbook(self.path)
        row = dict(zip((c.value for c in book['Pairs'][1]), (c.value for c in book['Pairs'][2])))
        self.assertEqual(row['repository'], 'game/example')
        self.assertEqual(book['Pairs']['C2'].data_type, 's')
        self.assertIsNone(row['manual_valid'])
        book.close()
        self.assertEqual(review.validated_rows(self.path), [])

    def test_noop_keeps_bytes_mtime_and_no_backup(self):
        self.sync()
        before, mtime = self.path.read_bytes(), self.path.stat().st_mtime_ns
        self.assertFalse(self.sync()['saved'])
        self.assertEqual(before, self.path.read_bytes())
        self.assertEqual(mtime, self.path.stat().st_mtime_ns)
        self.assertFalse((self.root / 'backups').exists())

    def test_append_preserves_reviews_custom_cells_sheet_style_and_backup(self):
        self.sync()
        self.accept()
        book = load_workbook(self.path)
        sheet = book['Pairs']
        sheet.cell(1, sheet.max_column + 1, 'my_comment')
        custom_column = sheet.max_column
        sheet.cell(2, custom_column, 'keep this')
        sheet.cell(2, custom_column).font = Font(bold=True, color='FF112233')
        sheet.cell(2, custom_column).comment = Comment('Evidence note', 'Ivan')
        book.create_sheet('My notes')['A1'] = '=1+1'
        book.save(self.path)
        book.close()
        before = self.path.read_bytes()
        self.write_run([self.b])  # old candidate absent from this run must remain
        self.assertEqual(self.sync()['added_pairs'], 1)
        book = load_workbook(self.path)
        self.assertEqual(book['Pairs'].max_row, 3)
        self.assertEqual(book['Pairs'].cell(2, custom_column).value, 'keep this')
        self.assertTrue(book['Pairs'].cell(2, custom_column).font.bold)
        self.assertEqual(book['Pairs'].cell(2, custom_column).comment.text, 'Evidence note')
        self.assertEqual(book['My notes']['A1'].value, '=1+1')
        book.close()
        self.assertEqual(len(review.validated_rows(self.path)), 1)
        self.assertEqual(next((self.root / 'backups').glob('*.xlsx')).read_bytes(), before)

    def test_duplicate_input_and_exact_url_variants(self):
        self.write_run([self.a, self.a, {**self.a, 'video_url': 'https://example.org/a.mp4?key=2#part'}])
        self.assertEqual(self.sync()['added_pairs'], 2)

    def test_duplicate_existing_record_aborts(self):
        self.sync()
        book = load_workbook(self.path)
        book['Pairs'].append([c.value for c in book['Pairs'][2]])
        book.save(self.path)
        book.close()
        before = self.path.read_bytes()
        with self.assertRaisesRegex(ValueError, 'duplicate record'):
            self.sync()
        self.assertEqual(before, self.path.read_bytes())

    def test_bad_identity_or_input_aborts_without_change(self):
        self.sync()
        before = self.path.read_bytes()
        self.write_run([{**self.a, 'issue_number': '-7'}])
        with self.assertRaises(ValueError):
            self.sync()
        self.assertEqual(before, self.path.read_bytes())

    def test_open_workbook_and_concurrent_sync_locks(self):
        self.sync()
        for lock in ('~$review.xlsx', '.review.xlsx.sync.lock', '.~lock.review.xlsx#'):
            with self.subTest(lock=lock):
                path = self.root / lock
                path.touch()
                with self.assertRaises(ValueError):
                    self.sync()
                path.unlink()

    def test_save_backup_replace_failures_keep_original(self):
        self.sync()
        self.write_run([self.a, self.b])
        for target in ('openpyxl.workbook.workbook.Workbook.save',
                       'scripts.review_workbook.shutil.copy2', 'scripts.review_workbook.os.replace'):
            with self.subTest(target=target):
                before = self.path.read_bytes()
                with patch(target, side_effect=OSError('injected failure')):
                    with self.assertRaises(OSError):
                        self.sync()
                self.assertEqual(before, self.path.read_bytes())
                self.assertFalse((self.root / '.review.xlsx.sync.lock').exists())

    def test_every_acceptance_prerequisite_is_required(self):
        self.sync()
        self.accept()
        self.assertEqual(len(review.validated_rows(self.path)), 1)
        for sheet_name, fields in (('Pairs', review.PAIR_CHECKS + ('bug_category',)),
                                   ('Repositories', review.REPO_CHECKS)):
            for field in fields:
                with self.subTest(field=field):
                    self.accept()
                    book = load_workbook(self.path)
                    self.change(book[sheet_name], 2, field, 'uncertain')
                    book.save(self.path)
                    book.close()
                    with self.assertRaisesRegex(ValueError, 'prerequisites'):
                        review.validated_rows(self.path)

    def test_export_and_unreviewed_rows(self):
        self.sync()
        self.accept()
        out = self.root / 'validated'
        self.assertEqual(review.export_validated(self.path, out), 1)
        with (out / 'validated_pairs.csv').open(encoding='utf-8-sig') as stream:
            row = next(csv.DictReader(stream))
        self.assertEqual(row['manual_valid'], 'yes')
        self.assertEqual(row['scope_valid'], 'yes')
        self.assertEqual(row['issue_title'], "'=untrusted text")
        raw = json.loads((out / 'validated_pairs.json').read_text())
        self.assertEqual(raw[0]['issue_title'], '=untrusted text')

    def test_csv_formula_prefixes_are_escaped(self):
        for text in ('=1+1', '+cmd', '-cmd', '@cmd', '  =1+1', '\tcmd', '\rcmd', '\ncmd'):
            with self.subTest(text=text):
                self.assertEqual(review.csv_text(text), "'" + text)
        self.assertEqual(review.csv_text('plain description'), 'plain description')
        self.assertEqual(review.csv_text(7), 7)

    def test_reordered_columns_and_sorted_rows(self):
        self.write_run([self.a, self.b])
        self.sync()
        self.accept()
        book = load_workbook(self.path)
        sheet = book['Pairs']
        values = list(sheet.values)
        for r, record in enumerate([values[0], values[2], values[1]], 1):
            for c, value in enumerate(reversed(record), 1):
                sheet.cell(r, c).value = value
        book.save(self.path)
        book.close()
        before = self.path.read_bytes()
        self.assertFalse(self.sync()['saved'])
        self.assertEqual(before, self.path.read_bytes())
        self.assertEqual(len(review.validated_rows(self.path)), 1)


if __name__ == '__main__':
    unittest.main()
