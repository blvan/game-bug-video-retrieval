"""Append candidates to a persistent Excel ledger; export fully reviewed pairs."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import os
import re
import shutil
import sys
import tempfile
from contextlib import contextmanager
from copy import copy
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit
from uuid import uuid4
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
CATEGORIES = ('rendering', 'animation', 'camera', 'spatial_collision', 'other_visual')
PAIR_CHECKS = ('video_accessible', 'is_gameplay', 'describes_bug', 'matches_report')
REPO_CHECKS = ('scope_valid', 'license_verified')
PAIR_FIELDS = [
    'repository', 'issue_number', 'issue_title', *PAIR_CHECKS, 'bug_category',
    'manual_valid', 'review_notes', 'reviewer', 'reviewed_at', 'issue_url', 'video_url',
    'candidate_id', 'candidate_score', 'media_status', 'repository_license',
    'source_url', 'issue_body', 'source_body', 'source_run', 'imported_at',
]
REPO_FIELDS = ['repository', *REPO_CHECKS, 'review_notes', 'reviewer', 'reviewed_at',
               'repository_url', 'detected_license', 'detected_scope', 'source_run']


def identity(row):
    repository = str(row.get('repository') or '').strip().lower()
    if not re.fullmatch(r'[a-z0-9_.-]+/[a-z0-9_.-]+', repository):
        raise ValueError(f'Invalid repository: {repository!r}')
    number = str(row.get('issue_number') or '').strip()
    if not number.isdecimal() or int(number) <= 0:
        raise ValueError(f'Invalid issue number: {number!r}')
    url = str(row.get('video_url') or '').strip()
    parsed = urlsplit(url)
    if parsed.scheme not in ('http', 'https') or not parsed.hostname:
        raise ValueError(f'Invalid video URL: {url!r}')
    # Keep URL query/fragment unchanged: removing them may change media identity.
    return repository, str(int(number)), url


def candidate_id(row):
    return hashlib.sha256(json.dumps(identity(row), ensure_ascii=False).encode()).hexdigest()[:24]


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        return list(csv.DictReader(stream))


def index_sheet(sheet, required, pair=False):
    headers = [cell.value for cell in sheet[1]]
    if len(headers) != len(set(headers)) or any(not h for h in headers):
        raise ValueError(f'{sheet.title}: blank or duplicate column headers')
    missing = set(required) - set(headers)
    if missing:
        raise ValueError(f'{sheet.title}: missing columns {sorted(missing)}')
    records = {}
    for cells in sheet.iter_rows(min_row=2):
        if all(c.value is None for c in cells):
            continue
        row = dict(zip(headers, (c.value for c in cells)))
        if pair:
            key = candidate_id(row)
            if row['candidate_id'] != key:
                raise ValueError(f'{sheet.title} row {cells[0].row}: identity/ID changed; restore original ID and links')
        else:
            key = str(row.get('repository') or '').strip().lower()
            if not re.fullmatch(r'[a-z0-9_.-]+/[a-z0-9_.-]+', key):
                raise ValueError(f'{sheet.title} row {cells[0].row}: invalid repository')
        if key in records:
            raise ValueError(f'{sheet.title}: duplicate record {key}')
        records[key] = row
    return headers, records


def write_literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'  # issue text must never become an Excel formula


def append_record(sheet, headers, record):
    if sheet.max_row >= 1048576:
        raise ValueError('Excel row limit reached')
    row_number = sheet.max_row + 1
    for column, field in enumerate(headers, 1):
        cell = sheet.cell(row_number, column)
        if row_number > 2:
            cell._style = copy(sheet.cell(row_number - 1, column)._style)
        else:
            cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(vertical='center', wrap_text=field in ('issue_title', 'review_notes'))
            if field in (*PAIR_CHECKS, *REPO_CHECKS, 'bug_category', 'manual_valid', 'review_notes', 'reviewer', 'reviewed_at'):
                cell.fill = PatternFill('solid', fgColor='FFF4D6')
        value = record.get(field, '')
        if isinstance(value, str) and len(value) > 32767:
            if field not in ('issue_body', 'source_body'):
                raise ValueError(f'{field}: value exceeds Excel cell limit')
            value = value[:32650] + '\n[Truncated in Excel; complete text remains in the source run.]'
        write_literal(cell, value)
        if field in ('issue_url', 'video_url', 'source_url', 'repository_url') and value:
            if urlsplit(str(value)).scheme in ('https', 'http'):
                cell.hyperlink = str(value)
        if isinstance(value, datetime):
            cell.number_format = 'yyyy-mm-dd hh:mm'
    title_lines = (len(str(record.get('issue_title', ''))) + 37) // 38
    sheet.row_dimensions[row_number].height = max(34, min(84, title_lines * 14 + 8))


def digest(path):
    return hashlib.sha256(path.read_bytes()).digest() if path.exists() else None


def check_supported(path):
    """Limit updates to our plain XLSX workflow; do not round-trip rich workbooks."""
    if path.suffix.lower() != '.xlsx':
        raise ValueError('Only plain .xlsx workbooks are supported')
    with ZipFile(path) as package:
        unsupported = [n for n in package.namelist() if any(p in n for p in (
            'xl/charts/', 'xl/pivot', 'xl/slicer', 'xl/externalLinks/', 'xl/activeX/',
            'xl/ctrlProps/', 'vbaProject', 'xl/threadedComments/', 'xl/richData/',
            'xl/media/', 'customXml/')) or (n.startswith('xl/drawings/') and not n.endswith('.vml'))]
        if unsupported:
            raise ValueError(f'Unsupported workbook features; original was not changed: {unsupported[0]}')


@contextmanager
def workbook_lock(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.with_name('~$' + path.name).exists() or path.with_name('.~lock.' + path.name + '#').exists():
        raise ValueError('Close the workbook in Excel/LibreOffice before syncing')
    lock = path.with_name('.' + path.name + '.sync.lock')
    try:
        fd = os.open(lock, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    except FileExistsError as error:
        raise ValueError(f'Another sync is active, or a stale lock needs inspection: {lock}') from error
    try:
        os.close(fd)
        yield
    finally:
        lock.unlink(missing_ok=True)


def sync(run_dirs, path, template=None):
    incoming, repos = {}, {}
    for run_dir in run_dirs:
        for row in read_csv(run_dir / 'video_issue_candidates.csv'):
            key = candidate_id(row)  # validate all input before opening the workbook
            repository, number, url = identity(row)
            incoming.setdefault(key, {**row, 'repository': repository, 'issue_number': int(number),
                                      'video_url': url, 'candidate_id': key, 'source_run': run_dir.name})
        metadata = run_dir / 'repositories.csv'
        if metadata.exists():
            for row in read_csv(metadata):
                repos.setdefault(row['repository'].strip().lower(), row)
    template = template or ROOT / 'templates' / 'review-template.xlsx'
    with workbook_lock(path):
        before = digest(path)
        check_supported(path if before is not None else template)
        book = load_workbook(path if before is not None else template)
        try:
            pairs, repositories = book['Pairs'], book['Repositories']
            pair_headers, old_pairs = index_sheet(pairs, PAIR_FIELDS, pair=True)
            repo_headers, old_repos = index_sheet(repositories, REPO_FIELDS)
            added_pairs = added_repos = 0
            now = datetime.now(timezone.utc).replace(tzinfo=None)
            for key, row in incoming.items():
                repository = row['repository']
                if repository not in old_repos:
                    meta = repos.get(repository, {})
                    record = {'repository': repository, 'repository_url': f'https://github.com/{repository}',
                              'detected_license': meta.get('license', row.get('repository_license', '')),
                              'detected_scope': meta.get('scope', ''), 'source_run': row['source_run']}
                    append_record(repositories, repo_headers, record)
                    old_repos[repository] = record
                    added_repos += 1
                if key not in old_pairs:
                    append_record(pairs, pair_headers, {**row, 'imported_at': now})
                    added_pairs += 1
            if not added_pairs and not added_repos and before is not None:
                return {'added_pairs': 0, 'added_repositories': 0, 'saved': False}
            for sheet in (pairs, repositories):
                ref = f'A1:{get_column_letter(sheet.max_column)}{sheet.max_row}'
                sheet.auto_filter.ref = ref
                for table in sheet.tables.values():
                    table.ref = ref
            fd, temporary = tempfile.mkstemp(prefix='.review-', suffix='.xlsx', dir=path.parent)
            os.close(fd)
            try:
                book.save(temporary)
                verification = load_workbook(temporary)
                try:
                    index_sheet(verification['Pairs'], PAIR_FIELDS, pair=True)
                    index_sheet(verification['Repositories'], REPO_FIELDS)
                finally:
                    verification.close()
                if digest(path) != before:
                    raise ValueError('Workbook changed while syncing; no replacement performed')
                if before is not None:
                    backup_dir = path.parent / 'backups'
                    backup_dir.mkdir(exist_ok=True)
                    stamp = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
                    backup = backup_dir / f'{path.stem}.{stamp}.{uuid4().hex[:6]}.xlsx'
                    shutil.copy2(path, backup)
                    if digest(backup) != before or digest(path) != before:
                        raise ValueError('Workbook changed during backup; no replacement performed')
                os.replace(temporary, path)
            finally:
                Path(temporary).unlink(missing_ok=True)
        finally:
            book.close()
    return {'added_pairs': added_pairs, 'added_repositories': added_repos, 'saved': True}


def validated_rows(path):
    check_supported(path)
    book = load_workbook(path, data_only=False)
    try:
        _, pairs = index_sheet(book['Pairs'], PAIR_FIELDS, pair=True)
        _, repos = index_sheet(book['Repositories'], REPO_FIELDS)
        accepted, invalid = [], []
        for key, row in pairs.items():
            decision = str(row.get('manual_valid') or '').strip().lower()
            if decision not in ('', 'yes', 'no', 'uncertain'):
                raise ValueError(f'{key}: invalid manual_valid value')
            if decision != 'yes':
                continue
            repository = repos.get(row['repository'].strip().lower(), {})
            missing = [f for f in PAIR_CHECKS if row.get(f) != 'yes']
            missing += [f for f in REPO_CHECKS if repository.get(f) != 'yes']
            if row.get('bug_category') not in CATEGORIES:
                missing.append('bug_category')
            if missing:
                invalid.append(f'{key}: {", ".join(missing)}')
            else:
                accepted.append({**row, **{f: repository[f] for f in REPO_CHECKS}})
        if invalid:
            raise ValueError('Rows marked yes have incomplete prerequisites: ' + '; '.join(invalid))
        return accepted
    finally:
        book.close()


def export_validated(path, output_dir):
    rows = validated_rows(path)  # validation must finish before replacing any exports
    output_dir.mkdir(parents=True, exist_ok=True)
    fields = PAIR_FIELDS + list(REPO_CHECKS)
    with (output_dir / 'validated_pairs.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction='ignore')
        writer.writeheader()
        writer.writerows({key: csv_text(value) for key, value in row.items()} for row in rows)
    (output_dir / 'validated_pairs.json').write_text(json.dumps(rows, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    return len(rows)


def csv_text(value):
    """Escape spreadsheet formula prefixes in CSV, leaving JSON text unchanged."""
    if isinstance(value, str) and (value.startswith(('\t', '\r', '\n')) or
                                   value.lstrip().startswith(('=', '+', '-', '@'))):
        return "'" + value
    return value


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    update = commands.add_parser('sync', help='Append previously unseen pairs; preserve all existing cells')
    update.add_argument('--run-dir', action='append', type=Path, required=True)
    update.add_argument('--workbook', type=Path, default=Path('data/review/review.xlsx'))
    export = commands.add_parser('export', help='Export only fully reviewed yes pairs')
    export.add_argument('--workbook', type=Path, default=Path('data/review/review.xlsx'))
    export.add_argument('--output-dir', type=Path, default=Path('data/validated'))
    args = parser.parse_args(argv)
    try:
        if args.command == 'sync':
            print(json.dumps(sync(args.run_dir, args.workbook)))
        else:
            print(json.dumps({'validated_pairs': export_validated(args.workbook, args.output_dir)}))
        return 0
    except (OSError, ValueError, KeyError) as error:
        logging.error('%s', error)
        return 1


if __name__ == '__main__':
    sys.exit(main())
